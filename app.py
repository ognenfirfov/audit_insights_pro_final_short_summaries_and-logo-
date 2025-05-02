import streamlit as st
from utils.processor import analyze_audits
import tempfile
import os
from docx import Document
from fpdf import FPDF

st.set_page_config(page_title="Audit Insights Pro+", layout="wide")
st.title("📘 Audit Insights Pro+")
st.markdown("Upload 2–3 audit reports in PDF format. You'll receive structured 1-page summaries, comparisons, and downloadable reports with a logo.")

uploaded_files = st.file_uploader("Upload PDF audit files", type=["pdf"], accept_multiple_files=True)

if uploaded_files:
    if len(uploaded_files) < 2 or len(uploaded_files) > 3:
        st.warning("Please upload exactly 2 or 3 audit reports.")
    else:
        with st.spinner("Analyzing files with AI..."):
            try:
                temp_paths, filenames = [], []
                for file in uploaded_files:
                    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
                    tmp.write(file.read())
                    temp_paths.append(tmp.name)
                    filenames.append(file.name)

                summaries, comparison, learnings, full_text = analyze_audits(temp_paths, filenames)

                st.markdown("## 🧾 Per-Audit Summaries")
                for i in range(len(summaries)):
                    with st.expander(f"Audit {i+1} Summary ({filenames[i]})"):
                        st.markdown(summaries[i])

                st.markdown("## 🟰 Comparison of Audits")
                st.markdown(comparison)

                st.markdown("## 📘 Learnings for Future Audits")
                st.markdown(learnings)

                # DOCX Export
                doc = Document()
                doc.add_heading("Audit Insights Pro Summary", 0)
                doc.add_paragraph(full_text)
                docx_path = os.path.join(tempfile.gettempdir(), "audit_summary.docx")
                doc.save(docx_path)
                with open(docx_path, "rb") as f:
                    st.download_button("📄 Download Summary (.docx)", f.read(), "audit_analysis_summary.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

                # PDF Export
                class PDF(FPDF):
                    def header(self):
                        logo = "utils/logo_audit_insights.png"
                        if os.path.exists(logo):
                            self.image(logo, 10, 8, 50)
                        self.set_font("Arial", "B", 14)
                        self.cell(0, 10, "Audit Insights Pro Summary", ln=True, align="C")
                        self.ln(10)
                    def chapter_body(self, text):
                        self.set_font("Arial", "", 11)
                        self.multi_cell(0, 10, text)

                pdf = PDF()
                pdf.add_page()
                pdf.chapter_body(full_text)
                pdf_path = os.path.join(tempfile.gettempdir(), "audit_summary.pdf")
                pdf.output(pdf_path)
                with open(pdf_path, "rb") as pdf_file:
                    st.download_button("📄 Download Summary (.pdf)", pdf_file.read(), "audit_analysis_summary.pdf", mime="application/pdf")
            except Exception as e:
                st.error(f"❌ An error occurred: {str(e)}")
else:
    st.info("⬆️ Please upload 2 or 3 PDF audit reports to begin.")
